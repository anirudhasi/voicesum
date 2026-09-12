import pytest
from unittest.mock import AsyncMock, patch, MagicMock
from docx import Document
import io

from routers.rom_router import download_final_docx

@pytest.mark.asyncio
async def test_download_final_docx_action_points_grouping_and_no_ids():
    """
    Verify that in download_final_docx:
    1. All agenda discussion points appear first in the table.
    2. Heading 'Action Points' appears after the agenda discussion points.
    3. All action points belonging to the agenda appear underneath the 'Action Points' heading.
    4. No point IDs (P1, P2, AP1, AP2) are displayed.
    5. No emojis (⚡) are displayed.
    """
    mock_rom_data = {
        "final_rom": {
            "include_action_points": True,
            "agendas": [
                {
                    "agenda_id": "A1",
                    "title": "Budget Ordinance Review",
                    "discussion_points": [
                        {
                            "id": "pt-1",
                            "text": "Rod Domowski presented Ordinance 2026-0083.",
                            "speaker": "Rod Domowski",
                            "action_points": [
                                {
                                    "task": "Rod Domowski assigned Staff to brief the council",
                                    "assignee": "Staff",
                                    "deadline": "2026-10-01"
                                }
                            ]
                        },
                        {
                            "id": "pt-2",
                            "text": "The committee reviewed options regarding the CHS loan repayment issue.",
                            "speaker": "Committee",
                            "action_points": [
                                {
                                    "task": "The Executive to transmit a report describing next steps",
                                    "assignee": "Executive",
                                    "deadline": "2026-10-15"
                                }
                            ]
                        }
                    ]
                }
            ]
        }
    }

    with patch("routers.rom_router._get_rom_data", AsyncMock(return_value=mock_rom_data)), \
         patch("routers.rom_router._validate_user_id", return_value="test-user"):

        response = await download_final_docx(
            recording_id="rec-123",
            version="long",
            include_action_points=True,
            current_user={"sub": "test-user"},
            db=MagicMock()
        )

        # Read the streamed DOCX into python-docx Document
        body_bytes = b""
        async for chunk in response.body_iterator:
            body_bytes += chunk

        doc = Document(io.BytesIO(body_bytes))
        assert len(doc.tables) >= 1
        table = doc.tables[0]

        # Check table contents
        rows_text = []
        for row in table.rows:
            row_cells_text = [c.text.strip() for c in row.cells]
            rows_text.append(row_cells_text)

        # Headers: ['ID', 'Agenda', 'Discussion Points', 'Action / Speaker']
        assert rows_text[0] == ['ID', 'Agenda', 'Discussion Points', 'Action / Speaker']

        # Find rows for discussion points and action points in the Discussion Points column (index 2)
        dp_col_values = [r[2] for r in rows_text[1:]]

        # Verify no point IDs like P1, P2, AP1, AP2 or emoji ⚡
        for r in rows_text:
            combined = " ".join(r)
            assert "P1:" not in combined and "P2:" not in combined and "P1" not in r[0]
            assert "AP1" not in combined and "AP2" not in combined
            assert "⚡" not in combined

        # Discussion point 1 must come first
        assert "Rod Domowski presented Ordinance 2026-0083." in dp_col_values[0]
        assert dp_col_values[0].startswith("•")

        # Discussion point 2 must come next
        assert "The committee reviewed options regarding the CHS loan repayment issue." in dp_col_values[1]
        assert dp_col_values[1].startswith("•")

        # Heading 'Action Points' must come after discussion points
        assert dp_col_values[2] == "Action Points"

        # Action points must follow the heading
        assert "Rod Domowski assigned Staff to brief the council" in dp_col_values[3]
        assert dp_col_values[3].startswith("•")
        assert "The Executive to transmit a report describing next steps" in dp_col_values[4]
        assert dp_col_values[4].startswith("•")


@pytest.mark.asyncio
async def test_download_final_docx_3column_single_default_agenda():
    """Verify 3-column table ordering and formatting for single default agenda."""
    mock_rom_data = {
        "final_rom": {
            "include_action_points": True,
            "agendas": [
                {
                    "agenda_id": "A1",
                    "title": "General Discussion",
                    "is_default_agenda": True,
                    "discussion_points": [
                        {
                            "id": "pt-1",
                            "text": "General topic discussed.",
                            "speaker": "Speaker 1",
                            "action_points": [
                                {
                                    "task": "Follow up on topic",
                                    "assignee": "Speaker 1",
                                    "deadline": None
                                }
                            ]
                        }
                    ]
                }
            ]
        }
    }

    with patch("routers.rom_router._get_rom_data", AsyncMock(return_value=mock_rom_data)), \
         patch("routers.rom_router._validate_user_id", return_value="test-user"):

        response = await download_final_docx(
            recording_id="rec-123",
            version=None,
            include_action_points=True,
            current_user={"sub": "test-user"},
            db=MagicMock()
        )

        body_bytes = b""
        async for chunk in response.body_iterator:
            body_bytes += chunk

        doc = Document(io.BytesIO(body_bytes))
        table = doc.tables[0]

        rows_text = [[c.text.strip() for c in r.cells] for r in table.rows]
        assert rows_text[0] == ['ID', 'Discussion Points', 'Action / Speaker']

        dp_col_values = [r[1] for r in rows_text[1:]]

        # No P1 or AP1 or ⚡
        for r in rows_text:
            combined = " ".join(r)
            assert "P1" not in combined and "AP1" not in combined and "⚡" not in combined

        # Discussion point first
        assert "General topic discussed." in dp_col_values[0]
        assert dp_col_values[0].startswith("•")

        # Heading 'Action Points'
        assert dp_col_values[1] == "Action Points"

        # Action point next
        assert "Follow up on topic" in dp_col_values[2]
        assert dp_col_values[2].startswith("•")
