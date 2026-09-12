"""Minutes of Meeting (MoM) router — generate, fetch, update, and export MoMs."""
import io
import json
import logging
import os
import re
import uuid
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import text

from database import get_db, get_db_context, dt_to_str, to_json, from_json
from routers.auth import get_current_user
from services.llm import generate_mom, analyze_writing_style, rewrite_section_content, summarize_long_points
from config import settings
from tasks.pipeline import _filter_high_confidence_segments
from xml.sax.saxutils import escape

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/mom", tags=["mom"])


class ActionItem(BaseModel):
    task: str = ""
    owner: str = ""
    deadline: str = ""


class MoMData(BaseModel):
    title: str = ""
    date: str = ""
    duration: float = 0.0
    planned_start_time: str = ""
    actual_start_time: str = ""
    planned_end_time: str = ""
    actual_end_time: str = ""
    participants: List[str] = []
    introduction: str = ""
    points_discussed: List[str] = []
    action_items: List[ActionItem] = []
    conclusion: str = ""


class RewriteSectionRequest(BaseModel):
    """Request body for the rewrite/section endpoint."""
    section: str = Field(..., description="'discussion_points' | 'action_items' | 'introduction' | 'conclusion'")
    content: Any = Field(..., description="List[str] for discussion_points/action_items, str for introduction/conclusion")
    rules: str = ""
    custom_prompt: str = ""
    window_size: int = Field(default=20, ge=1, le=200)


class RegenerateActionPointsRequest(BaseModel):
    """Request body for the regenerate-action-points endpoint."""
    window_minutes: float = Field(
        default=10.0, ge=1.0, le=60.0,
        description="Transcript window size in minutes for action point extraction (1–60, default 10).",
    )


class SummarizeLongPointsRequest(BaseModel):
    """Request body for the summarize-long-points endpoint."""
    threshold: int = Field(
        default=70, ge=10, le=1000,
        description="Word count threshold for long discussion points (default 70).",
    )


def _normalize_action_items(items: list) -> list:
    """Ensure each action item is a {task, owner, deadline} dict."""
    result = []
    vague_owners = {"my team", "our team", "you", "they", "everyone", "we", "someone", "somebody", "anybody", "team", "us"}
    for item in items:
        if isinstance(item, dict):
            owner_val = str(item.get("owner", "Unassigned") or "Unassigned").strip()
            if owner_val.lower() in vague_owners:
                owner_val = "Unassigned"
            result.append({
                "task": str(item.get("task", "") or ""),
                "owner": owner_val or "Unassigned",
                "deadline": str(item.get("deadline", "ASAP") or "ASAP"),
            })
        elif isinstance(item, str) and item.strip():
            result.append({"task": item.strip(), "owner": "Unassigned", "deadline": "ASAP"})
    return result


def _normalize_point_discussed(pt) -> str:
    """Convert a points_discussed entry (str or dict) to a clean 'Topic: Summary' string."""
    if isinstance(pt, str):
        s = pt.strip()
        if s.startswith('[') and s.endswith(']'):
            try:
                parsed = from_json(s, None)
                if isinstance(parsed, list):
                    return "\n".join(_normalize_point_discussed(item) for item in parsed if item)
            except Exception:
                pass
        return pt
    if isinstance(pt, dict):
        topic = str(pt.get('topic', '')).strip()
        summary = str(pt.get('summary') or pt.get('discussion_point') or pt.get('text', '')).strip()
        if topic and summary and not summary.lower().startswith(topic.lower()):
            return f"{topic}: {summary}"
        return summary or topic or str(pt)
    return str(pt)


def _mom_row_to_dict(mom) -> dict:
    """Convert a SQLite row for MoM into a clean dict with JSON fields deserialized."""
    raw_action_items = from_json(mom["action_items"], [])
    raw_points = from_json(mom["points_discussed"], []) or []
    raw_participants = from_json(mom["participants"], []) or []

    # Guard against double-encoded JSON strings
    if isinstance(raw_points, str):
        try:
            parsed_inner = from_json(raw_points, None)
            if isinstance(parsed_inner, list):
                raw_points = parsed_inner
        except Exception:
            pass

    if isinstance(raw_action_items, str):
        try:
            parsed_inner = from_json(raw_action_items, None)
            if isinstance(parsed_inner, list):
                raw_action_items = parsed_inner
        except Exception:
            pass

    if isinstance(raw_participants, str):
        try:
            parsed_inner = from_json(raw_participants, None)
            if isinstance(parsed_inner, list):
                raw_participants = parsed_inner
        except Exception:
            pass

    # Guard: if stored as a newline-separated string, split it
    if isinstance(raw_points, str):
        raw_points = [ln.strip() for ln in raw_points.splitlines() if ln.strip()] or [raw_points]

    clean_points = []
    for p in raw_points:
        norm = _normalize_point_discussed(p)
        if norm:
            # If norm contains newlines from an unpacked double-JSON array, flatten it
            if "\n" in norm and not isinstance(p, dict):
                clean_points.extend([sub.strip() for sub in norm.splitlines() if sub.strip()])
            else:
                clean_points.append(norm)

    return {
        "id": mom["id"],
        "recording_id": mom["recording_id"],
        "user_id": mom["user_id"],
        "title": mom.get("title") or "",
        "date": mom.get("date") or "",
        "duration": mom.get("duration") or 0,
        "planned_start_time": mom.get("planned_start_time") or "",
        "actual_start_time": mom.get("actual_start_time") or "",
        "planned_end_time": mom.get("planned_end_time") or "",
        "actual_end_time": mom.get("actual_end_time") or "",
        "participants": raw_participants if isinstance(raw_participants, list) else [],
        "introduction": mom.get("introduction") or "",
        "points_discussed": clean_points,
        "action_items": _normalize_action_items(raw_action_items),
        "conclusion": mom.get("conclusion") or "",
        "is_draft": bool(mom.get("is_draft", False)),
        "created_at": mom.get("created_at"),
        "updated_at": mom.get("updated_at"),
    }


@router.get("/{recording_id}")
async def get_mom(recording_id: str, current_user: dict = Depends(get_current_user)):
    user_id = current_user["id"]
    async with get_db_context() as db:
        r = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        mom = r.mappings().fetchone()

        r_spk = await db.execute(
            text("SELECT speaker_mappings FROM recordings WHERE id = :id AND user_id = :uid"),
            {"id": recording_id, "uid": user_id},
        )
        rec_spk = r_spk.mappings().fetchone()

    if not mom:
        raise HTTPException(status_code=404, detail="MoM not found")

    mom_dict = _mom_row_to_dict(mom)
    if rec_spk and rec_spk.get("speaker_mappings"):
        spk_mappings = from_json(rec_spk["speaker_mappings"], {})
        if spk_mappings and isinstance(spk_mappings, dict):
            from services.speaker_sync import apply_speaker_mappings_to_mom_dict
            mom_dict = apply_speaker_mappings_to_mom_dict(mom_dict, spk_mappings)

    return mom_dict


@router.post("/{recording_id}/generate")
async def generate_mom_endpoint(recording_id: str, current_user: dict = Depends(get_current_user)):
    user_id = current_user["id"]

    try:
        async with get_db_context() as db:
            r = await db.execute(
                text("""
                    SELECT transcript, raw_text, filename, created_at, duration,
                           speakers_detected, context_summary, context_summary_hash,
                           agenda_summary, reference_summary
                    FROM recordings WHERE id = :id AND user_id = :uid
                """),
                {"id": recording_id, "uid": user_id},
            )
            rec = r.mappings().fetchone()

        if not rec:
            raise HTTPException(status_code=404, detail="Recording not found")

        transcript = from_json(rec["transcript"], [])
        if not transcript:
            raise HTTPException(status_code=400, detail="No transcript available to summarize")

        # Filter low-confidence segments before generating MoM
        filtered_transcript = _filter_high_confidence_segments(transcript, settings.MIN_AVG_SEGMENT_CONFIDENCE)
        logger.info(
            f"[MoM] generate: using {len(filtered_transcript)}/{len(transcript)} high-confidence segments "
            f"for recording={recording_id}"
        )

        meta = {
            "filename": rec.get("filename", "Meeting Notes"),
            "created_at": rec.get("created_at", ""),
            "duration": rec.get("duration", 0),
            "speakers_detected": from_json(rec["speakers_detected"], []),
        }

        # Resolve context_summary (reuse cached or build fresh)
        import asyncio as _asyncio
        _loop = _asyncio.get_running_loop()
        from tasks.pipeline import _raw_text_hash
        from services.llm import build_context_summary as _build_ctx
        raw_text = rec.get("raw_text") or ""
        current_hash = _raw_text_hash(raw_text)
        ctx: str | None = None
        if rec.get("context_summary") and rec.get("context_summary_hash") == current_hash:
            ctx = rec["context_summary"]
            logger.info(f"[MoM] Reusing cached context_summary for {recording_id} ({len(ctx.split()):,} words)")
        else:
            logger.info(f"[MoM] Building fresh context_summary for {recording_id} ({len(raw_text)} chars)")
            try:
                ctx = await _loop.run_in_executor(None, _build_ctx, filtered_transcript)
                if ctx:
                    async with get_db_context() as db:
                        await db.execute(
                            text("UPDATE recordings SET context_summary = :ctx, context_summary_hash = :h "
                                 "WHERE id = :id AND user_id = :uid"),
                            {"ctx": ctx, "h": current_hash, "id": recording_id, "uid": user_id},
                        )
                        await db.commit()
            except Exception as ctx_err:
                logger.warning(f"[MoM] context_summary build failed (non-fatal): {ctx_err}")
                ctx = None

        # Retrieve stored agenda / reference summaries (uploaded via /attachments)
        agenda_summary: str | None = rec.get("agenda_summary") or None
        reference_summary: str | None = rec.get("reference_summary") or None
        if agenda_summary:
            logger.info(f"[MoM] Injecting agenda_summary ({len(agenda_summary)} chars) for {recording_id}")
        if reference_summary:
            logger.info(f"[MoM] Injecting reference_summary ({len(reference_summary)} chars) for {recording_id}")

        mom_data = await _loop.run_in_executor(
            None,
            lambda: generate_mom(
                filtered_transcript,
                meta,
                context=ctx or None,
                agenda_summary=agenda_summary,
                reference_summary=reference_summary,
            ),
        )

        now = datetime.now(timezone.utc)
        mom_id = str(uuid.uuid4())
        initial_version = [{"version": 1, "data": mom_data, "saved_at": dt_to_str(now)}]

        async with get_db_context() as db:
            # Check if MoM already exists (upsert pattern)
            r = await db.execute(
                text("SELECT id FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
                {"rid": recording_id, "uid": user_id},
            )
            existing = r.fetchone()

            if existing:
                await db.execute(
                    text("""
                        UPDATE minutes_of_meeting SET
                            title = :title, date = :date, duration = :duration,
                            planned_start_time = :planned_start_time,
                            actual_start_time = :actual_start_time,
                            participants = :participants,
                            introduction = :introduction,
                            points_discussed = :points_discussed,
                            action_items = :action_items,
                            conclusion = :conclusion,
                            versions = :versions, is_draft = 0, updated_at = :updated_at
                        WHERE recording_id = :rid AND user_id = :uid
                    """),
                    {
                        "title": mom_data.get("title", ""),
                        "date": mom_data.get("date", ""),
                        "duration": mom_data.get("duration", 0),
                        "planned_start_time": mom_data.get("planned_start_time", ""),
                        "actual_start_time": mom_data.get("actual_start_time", ""),
                        "planned_end_time": mom_data.get("planned_end_time", ""),
                        "actual_end_time": mom_data.get("actual_end_time", ""),
                        "participants": to_json(mom_data.get("participants", [])),
                        "introduction": mom_data.get("introduction", ""),
                        "points_discussed": to_json(mom_data.get("points_discussed", [])),
                        "action_items": to_json(mom_data.get("action_items", [])),
                        "conclusion": mom_data.get("conclusion", ""),
                        "versions": to_json(initial_version),
                        "updated_at": dt_to_str(now),
                        "rid": recording_id,
                        "uid": user_id,
                    },
                )
            else:
                await db.execute(
                    text("""
                        INSERT INTO minutes_of_meeting (
                            id, recording_id, user_id, title, date, duration,
                            planned_start_time, actual_start_time,
                            participants, introduction, points_discussed,
                            action_items, conclusion,
                            versions, is_draft, created_at, updated_at
                        )
                        VALUES (
                            :id, :rid, :uid, :title, :date, :duration,
                            :planned_start_time, :actual_start_time,
                            :participants, :introduction, :points_discussed,
                            :action_items, :conclusion,
                            :versions, 0, :created_at, :updated_at
                        )
                    """),
                    {
                        "id": mom_id,
                        "rid": recording_id,
                        "uid": user_id,
                        "title": mom_data.get("title", ""),
                        "date": mom_data.get("date", ""),
                        "duration": mom_data.get("duration", 0),
                        "planned_start_time": mom_data.get("planned_start_time", ""),
                        "actual_start_time": mom_data.get("actual_start_time", ""),
                        "planned_end_time": mom_data.get("planned_end_time", ""),
                        "actual_end_time": mom_data.get("actual_end_time", ""),
                        "participants": to_json(mom_data.get("participants", [])),
                        "introduction": mom_data.get("introduction", ""),
                        "points_discussed": to_json(mom_data.get("points_discussed", [])),
                        "action_items": to_json(mom_data.get("action_items", [])),
                        "conclusion": mom_data.get("conclusion", ""),
                        "versions": to_json(initial_version),
                        "created_at": dt_to_str(now),
                        "updated_at": dt_to_str(now),
                    },
                )
            await db.commit()

            # Fetch the saved record
            r2 = await db.execute(
                text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
                {"rid": recording_id, "uid": user_id},
            )
            saved = r2.mappings().fetchone()

        return _mom_row_to_dict(saved)
    finally:
        from services.ai_provider import QwenProvider
        QwenProvider.unload_model()



@router.patch("/{recording_id}")
async def update_mom(recording_id: str, data: MoMData, current_user: dict = Depends(get_current_user)):
    user_id = current_user["id"]

    async with get_db_context() as db:
        r = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        mom = r.mappings().fetchone()
        if not mom:
            raise HTTPException(status_code=404, detail="MoM not found")

        now = datetime.now(timezone.utc)
        update_data = data.dict()

        # Version tracking: push a version if last version is older than 5 min
        versions = from_json(mom["versions"], [])
        last_version = versions[-1] if versions else None
        push_version = False
        if last_version:
            last_saved_str = last_version.get("saved_at")
            if last_saved_str:
                try:
                    last_saved = datetime.fromisoformat(last_saved_str)
                    if last_saved.tzinfo is None:
                        last_saved = last_saved.replace(tzinfo=timezone.utc)
                    if (now - last_saved).total_seconds() > 300:
                        push_version = True
                except Exception:
                    pass

        if push_version:
            new_version = {
                "version": len(versions) + 1,
                "data": update_data,
                "saved_at": dt_to_str(now),
            }
            versions.append(new_version)

        await db.execute(
            text("""
                UPDATE minutes_of_meeting SET
                    title = :title, date = :date, duration = :duration,
                    planned_start_time = :planned_start_time,
                    actual_start_time = :actual_start_time,
                    planned_end_time = :planned_end_time,
                    actual_end_time = :actual_end_time,
                    participants = :participants,
                    introduction = :introduction,
                    points_discussed = :points_discussed,
                    action_items = :action_items,
                    conclusion = :conclusion,
                    versions = :versions, is_draft = 1, updated_at = :updated_at
                WHERE recording_id = :rid AND user_id = :uid
            """),
            {
                "title": update_data.get("title", ""),
                "date": update_data.get("date", ""),
                "duration": update_data.get("duration", 0),
                "planned_start_time": update_data.get("planned_start_time", ""),
                "actual_start_time": update_data.get("actual_start_time", ""),
                "planned_end_time": update_data.get("planned_end_time", ""),
                "actual_end_time": update_data.get("actual_end_time", ""),
                "participants": to_json(update_data.get("participants", [])),
                "introduction": update_data.get("introduction", ""),
                "points_discussed": to_json(update_data.get("points_discussed", [])),
                "action_items": to_json(update_data.get("action_items", [])),
                "conclusion": update_data.get("conclusion", ""),
                "versions": to_json(versions),
                "updated_at": dt_to_str(now),
                "rid": recording_id,
                "uid": user_id,
            },
        )
        await db.commit()

    return {"status": "success", "updated_at": dt_to_str(now)}


@router.post("/{recording_id}/version")
async def save_mom_version(
    recording_id: str,
    data: Optional[MoMData] = None,
    current_user: dict = Depends(get_current_user),
):
    """
    Explicitly save the current MoM as a new named version entry in history.
    """
    user_id = current_user["id"]

    async with get_db_context() as db:
        r = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        mom = r.mappings().fetchone()
        if not mom:
            raise HTTPException(status_code=404, detail="MoM not found")

        now = datetime.now(timezone.utc)
        if data:
            update_data = data.dict()
        else:
            update_data = _mom_row_to_dict(mom)

        versions = from_json(mom["versions"], [])
        new_version_num = len(versions) + 1
        new_version = {
            "version": new_version_num,
            "data": {
                "title": update_data.get("title", ""),
                "date": update_data.get("date", ""),
                "duration": update_data.get("duration", 0),
                "planned_start_time": update_data.get("planned_start_time", ""),
                "actual_start_time": update_data.get("actual_start_time", ""),
                "planned_end_time": update_data.get("planned_end_time", ""),
                "actual_end_time": update_data.get("actual_end_time", ""),
                "participants": update_data.get("participants", []),
                "introduction": update_data.get("introduction", ""),
                "points_discussed": update_data.get("points_discussed", []),
                "action_items": update_data.get("action_items", []),
                "conclusion": update_data.get("conclusion", ""),
            },
            "saved_at": dt_to_str(now),
            "label": "Manual Save",
        }
        versions.append(new_version)

        await db.execute(
            text("""
                UPDATE minutes_of_meeting SET
                    title = :title, date = :date, duration = :duration,
                    planned_start_time = :planned_start_time,
                    actual_start_time = :actual_start_time,
                    planned_end_time = :planned_end_time,
                    actual_end_time = :actual_end_time,
                    participants = :participants,
                    introduction = :introduction,
                    points_discussed = :points_discussed,
                    action_items = :action_items,
                    conclusion = :conclusion,
                    versions = :versions, is_draft = 1, updated_at = :updated_at
                WHERE recording_id = :rid AND user_id = :uid
            """),
            {
                "title": update_data.get("title", ""),
                "date": update_data.get("date", ""),
                "duration": update_data.get("duration", 0),
                "planned_start_time": update_data.get("planned_start_time", ""),
                "actual_start_time": update_data.get("actual_start_time", ""),
                "planned_end_time": update_data.get("planned_end_time", ""),
                "actual_end_time": update_data.get("actual_end_time", ""),
                "participants": to_json(update_data.get("participants", [])),
                "introduction": update_data.get("introduction", ""),
                "points_discussed": to_json(update_data.get("points_discussed", [])),
                "action_items": to_json(_normalize_action_items(update_data.get("action_items", []))),
                "conclusion": update_data.get("conclusion", ""),
                "versions": to_json(versions),
                "updated_at": dt_to_str(now),
                "rid": recording_id,
                "uid": user_id,
            },
        )
        await db.commit()

    return {"status": "success", "version_number": new_version_num, "versions": versions}


@router.get("/{recording_id}/versions")
async def get_mom_versions(recording_id: str, current_user: dict = Depends(get_current_user)):
    user_id = current_user["id"]
    async with get_db_context() as db:
        r = await db.execute(
            text("SELECT versions FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        row = r.mappings().fetchone()

    if not row:
        raise HTTPException(status_code=404, detail="MoM not found")

    return {"versions": from_json(row["versions"], [])}


@router.post("/{recording_id}/docx")
@router.post("/{recording_id}/pdf")
async def export_mom_docx(
    recording_id: str,
    data: Optional[MoMData] = None,
    current_user: dict = Depends(get_current_user)
):
    user_id = current_user["id"]

    # Check if the input data is missing or effectively empty
    is_empty = True
    if data:
        d = data.dict()
        # If any of the main content fields are present, we consider it non-empty
        if (d.get("title") or d.get("introduction") or d.get("conclusion") or
            d.get("points_discussed") or d.get("action_items") or d.get("participants")):
            is_empty = False

    if is_empty:
        # Load MoM from database (which also verifies ownership)
        async with get_db_context() as db:
            r = await db.execute(
                text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
                {"rid": recording_id, "uid": user_id},
            )
            mom_row = r.mappings().fetchone()

        if not mom_row:
            raise HTTPException(status_code=404, detail="MoM not found")
        mom = _mom_row_to_dict(mom_row)
    else:
        # Verify ownership / existence
        async with get_db_context() as db:
            r = await db.execute(
                text("SELECT id FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
                {"rid": recording_id, "uid": user_id},
            )
            exists = r.fetchone()

        if not exists:
            raise HTTPException(status_code=404, detail="MoM not found")

        # Use the caller-supplied data (live editor state)
        mom = data.dict()

    mom["action_items"] = _normalize_action_items(mom.get("action_items", []))

    try:
        docx_bytes = generate_mom_docx(mom)
    except Exception as e:
        logger.error(f"Failed to build MoM DOCX: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"DOCX generation failed: {e}")

    safe_title = str(mom.get("title") or "Meeting").replace(" ", "_").replace("/", "-")
    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="MoM_{safe_title}.docx"'},
    )


def generate_mom_docx(mom: dict) -> bytes:
    """Generate a formatted Microsoft Word (.docx) document matching MoM layout."""
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = docx.Document()

    # Set page margins to 0.7 inches
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Base font style
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Arial'
    normal_font.size = Pt(10)
    normal_font.color.rgb = RGBColor(51, 65, 85)

    def add_hr(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr %s><w:bottom w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/></w:pBdr>' % nsdecls('w'))
        pPr.append(pBdr)

    # 1. Title Block
    title_text = str(mom.get("title") or "Minutes of Meeting").strip()
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title_text)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)
    add_hr(p_title)

    # 2. Meeting Meta
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(6)
    p_meta.paragraph_format.space_after = Pt(12)
    p_meta.paragraph_format.line_spacing = 1.25

    meta_fields = []
    meta_fields.append(("Date", str(mom.get("date") or "Unknown")))
    members = ", ".join(mom.get("participants", []))
    meta_fields.append(("Members", members if members else "N/A"))
    if mom.get("planned_start_time"):
        meta_fields.append(("Planned Starting Time", str(mom["planned_start_time"])))
    if mom.get("actual_start_time"):
        meta_fields.append(("Actual Starting Time", str(mom["actual_start_time"])))
    if mom.get("planned_end_time"):
        meta_fields.append(("Planned End Time", str(mom["planned_end_time"])))
    if mom.get("actual_end_time"):
        meta_fields.append(("Actual End Time", str(mom["actual_end_time"])))

    for idx, (lbl, val) in enumerate(meta_fields):
        r_lbl = p_meta.add_run(f"{lbl}: ")
        r_lbl.bold = True
        r_lbl.font.color.rgb = RGBColor(51, 65, 85)
        r_val = p_meta.add_run(f"{val}" + ("" if idx == len(meta_fields) - 1 else "\n"))
        r_val.font.color.rgb = RGBColor(51, 65, 85)

    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
        add_hr(p)

    # 3. Introduction
    if mom.get("introduction"):
        add_section_heading("INTRODUCTION")
        p_intro = doc.add_paragraph()
        p_intro.paragraph_format.space_before = Pt(4)
        p_intro.paragraph_format.space_after = Pt(8)
        p_intro.paragraph_format.line_spacing = 1.2
        r_intro = p_intro.add_run(mom["introduction"])
        r_intro.font.size = Pt(10)
        r_intro.font.color.rgb = RGBColor(51, 65, 85)

    # 4. Points Discussed
    if mom.get("points_discussed"):
        add_section_heading("POINTS DISCUSSED")
        for pt in mom["points_discussed"]:
            p_pt = doc.add_paragraph(style='List Bullet')
            p_pt.paragraph_format.space_before = Pt(2)
            p_pt.paragraph_format.space_after = Pt(4)
            p_pt.paragraph_format.line_spacing = 1.15
            p_pt.paragraph_format.left_indent = Inches(0.25)
            r_pt = p_pt.add_run(str(pt))
            r_pt.font.size = Pt(10)
            r_pt.font.color.rgb = RGBColor(51, 65, 85)

    # 5. Action Points
    if mom.get("action_items"):
        add_section_heading("ACTION POINTS")

        from collections import defaultdict
        by_owner: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for ai in mom["action_items"]:
            owner = ai.get("owner", "Unassigned") if isinstance(ai, dict) else "Unassigned"
            by_owner[owner].append(ai)

        general = by_owner.pop("Unassigned", [])

        for owner, owner_items in sorted(by_owner.items()):
            p_owner = doc.add_paragraph()
            p_owner.paragraph_format.space_before = Pt(8)
            p_owner.paragraph_format.space_after = Pt(2)
            r_owner = p_owner.add_run(str(owner))
            r_owner.font.bold = True
            r_owner.font.size = Pt(10.5)
            r_owner.font.color.rgb = RGBColor(30, 41, 59)

            for ai in owner_items:
                t = ai.get('task', '') if isinstance(ai, dict) else str(ai)
                d = ai.get('deadline', '') if isinstance(ai, dict) else ''
                p_ai = doc.add_paragraph(style='List Bullet')
                p_ai.paragraph_format.space_before = Pt(1)
                p_ai.paragraph_format.space_after = Pt(3)
                p_ai.paragraph_format.left_indent = Inches(0.4)
                r_task = p_ai.add_run(t)
                r_task.font.size = Pt(10)
                r_task.font.color.rgb = RGBColor(51, 65, 85)
                if d != "" and (d.lower() not in ("asap", "none", "null", "")):
                    r_due = p_ai.add_run(f" (Due: {d})")
                    r_due.font.italic = True
                    r_due.font.size = Pt(9.5)
                    r_due.font.color.rgb = RGBColor(100, 116, 139)

        if general:
            p_gen = doc.add_paragraph()
            p_gen.paragraph_format.space_before = Pt(8)
            p_gen.paragraph_format.space_after = Pt(2)
            r_gen = p_gen.add_run("General")
            r_gen.font.bold = True
            r_gen.font.size = Pt(10.5)
            r_gen.font.color.rgb = RGBColor(30, 41, 59)

            for ai in general:
                t = ai.get('task', '') if isinstance(ai, dict) else str(ai)
                d = ai.get('deadline', '') if isinstance(ai, dict) else ''
                p_ai = doc.add_paragraph(style='List Bullet')
                p_ai.paragraph_format.space_before = Pt(1)
                p_ai.paragraph_format.space_after = Pt(3)
                p_ai.paragraph_format.left_indent = Inches(0.4)
                r_task = p_ai.add_run(t)
                r_task.font.size = Pt(10)
                r_task.font.color.rgb = RGBColor(51, 65, 85)
                if d != "" and (d.lower() not in ("asap", "none", "null", "")):
                    r_due = p_ai.add_run(f" (Due: {d})")
                    r_due.font.italic = True
                    r_due.font.size = Pt(9.5)
                    r_due.font.color.rgb = RGBColor(100, 116, 139)

    # 6. Conclusion
    if mom.get("conclusion"):
        add_section_heading("CONCLUSION")
        p_conc = doc.add_paragraph()
        p_conc.paragraph_format.space_before = Pt(4)
        p_conc.paragraph_format.space_after = Pt(8)
        p_conc.paragraph_format.line_spacing = 1.2
        r_conc = p_conc.add_run(mom["conclusion"])
        r_conc.font.size = Pt(10)
        r_conc.font.color.rgb = RGBColor(51, 65, 85)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()



@router.get("/{recording_id}/docx")
async def download_mom_docx(
    recording_id: str,
    current_user: dict = Depends(get_current_user),
    db=Depends(get_db)
):
    """Download the main MoM as a Word (.docx) document."""
    user_id = current_user["id"]
    async with get_db_context() as db_session:
        r_rec = await db_session.execute(
            text("SELECT filename, speaker_mappings FROM recordings WHERE id = :id AND user_id = :uid"),
            {"id": recording_id, "uid": user_id},
        )
        rec_row = r_rec.mappings().fetchone()

        r_mom = await db_session.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        mom_row = r_mom.mappings().fetchone()

    if not mom_row:
        raise HTTPException(status_code=404, detail="MoM not found")

    mom = _mom_row_to_dict(mom_row)
    if rec_row and rec_row.get("speaker_mappings"):
        spk_mappings = from_json(rec_row["speaker_mappings"], {})
        if spk_mappings and isinstance(spk_mappings, dict):
            from services.speaker_sync import apply_speaker_mappings_to_mom_dict
            mom = apply_speaker_mappings_to_mom_dict(mom, spk_mappings)

    from docx import Document
    doc = Document()

    title = mom.get("title") or (rec_row.get("filename") if rec_row else "Minutes of Meeting")
    h = doc.add_heading(title, level=0)
    for r in h.runs:
        r.bold = True

    if mom.get("date"):
        p = doc.add_paragraph()
        run = p.add_run("Date: ")
        run.bold = True
        p.add_run(str(mom["date"]))

    if mom.get("duration"):
        p = doc.add_paragraph()
        run = p.add_run("Duration: ")
        run.bold = True
        p.add_run(str(mom["duration"]))

    participants = mom.get("participants", [])
    if participants:
        p = doc.add_paragraph()
        run = p.add_run("Participants: ")
        run.bold = True
        p.add_run(", ".join(str(part) for part in participants))

    if mom.get("introduction"):
        h1 = doc.add_heading("1. Introduction & Overview", level=1)
        for r in h1.runs: r.bold = True
        doc.add_paragraph(mom["introduction"])

    pts = mom.get("points_discussed", [])
    if pts:
        h2 = doc.add_heading("2. Key Discussion Points", level=1)
        for r in h2.runs: r.bold = True
        for idx, pt in enumerate(pts, 1):
            if isinstance(pt, dict):
                h_sub = doc.add_heading(f"2.{idx} {pt.get('topic', 'Discussion')}", level=2)
                for r in h_sub.runs: r.bold = True
                if pt.get("summary"):
                    doc.add_paragraph(pt["summary"])
            else:
                doc.add_paragraph(f"• {str(pt)}")

    actions = mom.get("action_items", [])
    if actions:
        h3 = doc.add_heading("3. Action Items", level=1)
        for r in h3.runs: r.bold = True
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = 'Table Grid'
        hdr = tbl.rows[0].cells
        hdr[0].text = "#"
        hdr[1].text = "Action Item"
        hdr[2].text = "Owner"
        hdr[3].text = "Deadline"
        for cell in hdr:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

        for idx, act in enumerate(actions, 1):
            r_cells = tbl.add_row().cells
            r_cells[0].text = str(idx)
            r_cells[1].text = act.get("task") or act.get("item") or act.get("description") or "-"
            r_cells[2].text = act.get("owner") or "Unassigned"
            r_cells[3].text = act.get("deadline") or "ASAP"

    if mom.get("conclusion"):
        h4 = doc.add_heading("4. Conclusion", level=1)
        for r in h4.runs: r.bold = True
        doc.add_paragraph(mom["conclusion"])

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    safe_title = str(mom.get("title") or "Meeting").replace(" ", "_").replace("/", "-")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="MoM_{safe_title}.docx"'},
    )


# ══════════════════════════════════════════════════════════════
# Rewrite Endpoints
# ══════════════════════════════════════════════════════════════


MAX_REWRITE_DOC_SIZE_MB = 30
ALLOWED_REWRITE_EXTENSIONS = {
    ".pdf", ".docx", ".pptx", ".txt", ".md",
    ".png", ".jpg", ".jpeg", ".webp",
    ".xlsx", ".xls", ".csv",
}


@router.post("/{recording_id}/rewrite/analyze-style")
async def rewrite_analyze_style(
    recording_id: str,
    section: str = Form(..., description="'discussion_points' or 'action_items'"),
    files: List[UploadFile] = File(...),
    current_user: dict = Depends(get_current_user),
):
    """
    Upload one or more reference documents and extract writing style rules
    that describe how discussion_points or action_items are written in those docs.
    Returns a newline-separated list of style rules as a string.
    """
    user_id = current_user["id"]

    if section not in ("discussion_points", "action_items"):
        raise HTTPException(status_code=400, detail="section must be 'discussion_points' or 'action_items'")

    # Verify MoM ownership
    async with get_db_context() as db:
        r = await db.execute(
            text("SELECT id FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        if not r.fetchone():
            raise HTTPException(status_code=404, detail="MoM not found")

    # Read and validate uploaded files
    import tempfile
    from services.doc_extractor import extract_text_from_file

    texts = []
    tmp_paths = []
    try:
        for upload in files:
            filename = upload.filename or "upload"
            ext = os.path.splitext(filename.lower())[1]
            if ext not in ALLOWED_REWRITE_EXTENSIONS:
                raise HTTPException(
                    status_code=400,
                    detail=f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(ALLOWED_REWRITE_EXTENSIONS))}",
                )

            data = await upload.read()
            if len(data) > MAX_REWRITE_DOC_SIZE_MB * 1024 * 1024:
                raise HTTPException(status_code=413, detail=f"File '{filename}' exceeds {MAX_REWRITE_DOC_SIZE_MB} MB limit.")

            # Write to a temp file so doc_extractor can open it
            suffix = ext
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                tmp.write(data)
                tmp_path = tmp.name
            tmp_paths.append(tmp_path)

            text_content = extract_text_from_file(tmp_path, filename)
            if text_content.strip():
                texts.append(f"=== {filename} ===\n{text_content.strip()}")
            else:
                logger.warning(f"[MoM Rewrite] No text extracted from '{filename}'")
    finally:
        for p in tmp_paths:
            try:
                os.remove(p)
            except Exception:
                pass

    if not texts:
        raise HTTPException(
            status_code=422,
            detail="Could not extract any readable text from the uploaded files.",
        )

    combined_text = "\n\n".join(texts)
    logger.info(f"[MoM Rewrite] Analyzing style for section='{section}', {len(texts)} file(s), {len(combined_text)} chars")

    import asyncio as _asyncio
    _loop = _asyncio.get_running_loop()

    try:
        rules = await _loop.run_in_executor(
            None,
            lambda: analyze_writing_style(combined_text, section),
        )
    except Exception as e:
        logger.error(f"[MoM Rewrite] Style analysis failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Style analysis failed: {e}")
    finally:
        from services.ai_provider import QwenProvider
        QwenProvider.unload_model()

    if not rules or not rules.strip():
        raise HTTPException(status_code=500, detail="LLM returned empty style rules. Please try again.")

    return {"rules": rules.strip()}


@router.post("/{recording_id}/regenerate-action-points")
async def regenerate_mom_action_points(
    recording_id: str,
    req: RegenerateActionPointsRequest,
    current_user: dict = Depends(get_current_user),
):
    """Regenerate MoM action points from the original transcript using a dedicated LLM extraction call.

    Divides the meeting transcript into time windows of `window_minutes` minutes, runs the
    MOM_REGENERATE_ACTION_POINTS_PROMPT on each window concurrently, merges all extracted
    action items into a flat list, maps them to the MoM schema, and saves the result as a
    new MoM version. The rest of the MoM (introduction, discussion points, conclusion) is
    untouched.
    """
    import asyncio as _asyncio
    from concurrent.futures import ThreadPoolExecutor
    from services.ai_provider import get_provider
    from tasks.pipeline import _filter_high_confidence_segments

    user_id = current_user["id"]
    window_secs = req.window_minutes * 60.0

    async with get_db_context() as db:
        # Load transcript from recordings table
        r_rec = await db.execute(
            text("SELECT transcript, status FROM recordings WHERE id = :id AND user_id = :uid"),
            {"id": recording_id, "uid": user_id},
        )
        rec = r_rec.mappings().fetchone()
        if not rec:
            raise HTTPException(status_code=404, detail="Recording not found.")
        if rec.get("status") not in ("done", "error", "transcript_ready"):
            raise HTTPException(
                status_code=400,
                detail="Recording is still being processed. Wait until transcription finishes.",
            )

        # Load the current MoM row (needed for unchanged fields + version history)
        r_mom = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        mom_row = r_mom.mappings().fetchone()

        # Load user settings for parallel concurrency
        r_us = await db.execute(
            text("SELECT rom_parallel_window_processing FROM user_settings WHERE user_id = :uid"),
            {"uid": user_id},
        )
        us_row = r_us.fetchone()
        concurrency = us_row[0] if us_row and us_row[0] is not None else getattr(settings, "ROM_PARALLEL_WINDOW_PROCESSING", 2)
        try:
            concurrency = max(1, min(5, int(concurrency)))
        except (ValueError, TypeError):
            concurrency = 2

    if not mom_row:
        raise HTTPException(status_code=404, detail="MoM not found. Generate the MoM first.")

    transcript = from_json(rec["transcript"], [])
    if not transcript:
        raise HTTPException(status_code=400, detail="No transcript found for this recording.")

    transcript = _filter_high_confidence_segments(transcript, settings.MIN_AVG_SEGMENT_CONFIDENCE)
    if not transcript:
        raise HTTPException(status_code=400, detail="No high-confidence transcript segments available.")

    logger.info(
        f"[MoM Regen Actions] recording_id={recording_id}, "
        f"{len(transcript)} segments, window={req.window_minutes}min, concurrency={concurrency}"
    )

    # ── Build time windows ────────────────────────────────────────────────────
    def _build_windows(segs, win_secs):
        """Group transcript segments into sequential time windows."""
        windows, current = [], []
        win_start = segs[0].get("start", 0.0) if segs else 0.0
        for seg in segs:
            seg_end = seg.get("end", 0.0)
            if seg_end - win_start > win_secs and current:
                windows.append(current)
                current = [seg]
                win_start = seg.get("start", 0.0)
            else:
                current.append(seg)
        if current:
            windows.append(current)
        return windows

    windows = _build_windows(transcript, window_secs)
    total_windows = len(windows)
    logger.info(f"[MoM Regen Actions] {total_windows} window(s) to process.")

    # ── Per-window extraction (parallel concurrency from settings) ──────────
    def _format_window(segs):
        lines = []
        for seg in segs:
            spk = seg.get("speaker", "Unknown")
            s = seg.get("start", 0.0)
            e = seg.get("end", 0.0)
            txt = seg.get("text", "").strip()
            lines.append(f"[{s:.1f}-{e:.1f}] {spk}: {txt}")
        return "\n".join(lines)

    provider = get_provider()

    def _process_window(idx_win):
        idx, win = idx_win
        wt = _format_window(win)
        try:
            res = provider.regenerate_mom_action_points(wt)
            items = res.get("action_items") or []
            logger.info(f"[MoM Regen Actions] Window {idx+1}/{total_windows}: {len(items)} item(s)")
            return items
        except Exception as we:
            logger.warning(f"[MoM Regen Actions] Window {idx+1} failed: {we}")
            return []

    loop = _asyncio.get_running_loop()
    try:
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            raw_results = await loop.run_in_executor(
                None,
                lambda: list(pool.map(_process_window, enumerate(windows))),
            )
    finally:
        try:
            from services.ai_provider import QwenProvider
            QwenProvider.unload_model()
        except Exception:
            pass

    # ── Map LLM output → MoM ActionItem schema ───────────────────────────────
    def _map_action_item(raw):
        """Convert {assigner, assignee, task, deadline} → {task, owner, deadline}."""
        if not isinstance(raw, dict):
            return None
        task_text = str(raw.get("task") or "").strip()
        if not task_text:
            return None
        assigner = str(raw.get("assigner") or "").strip() or None
        null_words = {"null", "none", "unassigned", "n/a", ""}
        if assigner and assigner.lower() in null_words:
            assigner = None
        # Build display task: include assigner attribution if present
        if assigner:
            display_task = f"{task_text} (Assigned by: {assigner})"
        else:
            display_task = task_text
        assignee = str(raw.get("assignee") or "").strip() or None
        if assignee and assignee.lower() in null_words:
            assignee = None
        deadline = str(raw.get("deadline") or "").strip() or None
        if deadline and deadline.lower() in null_words:
            deadline = None
        return {
            "task": display_task,
            "owner": assignee or "Unassigned",
            "deadline": deadline or "ASAP",
        }

    # ── Collect raw extracted action items across all windows ─────────────────
    all_raw_action_items = []
    for window_items in raw_results:
        if isinstance(window_items, list):
            for raw_item in window_items:
                if isinstance(raw_item, dict) and str(raw_item.get("task") or "").strip():
                    all_raw_action_items.append(raw_item)

    logger.info(f"[MoM Regen Actions] Total raw action items extracted across {total_windows} windows: {len(all_raw_action_items)}")

    # ── Final global deduplication pass ───────────────────────────────────────
    if all_raw_action_items:
        try:
            logger.info(f"[MoM Regen Actions] Running final global deduplication pass on {len(all_raw_action_items)} action items...")
            deduped_raw_items = provider.deduplicate_action_points(all_raw_action_items)
            if isinstance(deduped_raw_items, list) and len(deduped_raw_items) > 0:
                logger.info(f"[MoM Regen Actions] Global deduplication complete: {len(all_raw_action_items)} -> {len(deduped_raw_items)} item(s)")
                all_raw_action_items = deduped_raw_items
            else:
                logger.warning("[MoM Regen Actions] Global deduplication returned empty or invalid list, keeping extracted items.")
        except Exception as de:
            logger.warning(f"[MoM Regen Actions] Global deduplication failed, keeping extracted items: {de}")

    # ── Map LLM output → MoM ActionItem schema ───────────────────────────────
    all_items = []
    for raw_item in all_raw_action_items:
        mapped = _map_action_item(raw_item)
        if mapped:
            all_items.append(mapped)

    logger.info(f"[MoM Regen Actions] Total final action items after deduplication & mapping: {len(all_items)}")

    # ── Persist to DB with a new version entry ────────────────────────────────
    mom_dict = _mom_row_to_dict(mom_row)
    mom_dict["action_items"] = all_items  # replace action_items only

    now = datetime.now(timezone.utc)
    versions = from_json(mom_row["versions"], [])
    new_version_num = len(versions) + 1
    new_version = {
        "version": new_version_num,
        "data": {
            "title": mom_dict.get("title", ""),
            "date": mom_dict.get("date", ""),
            "duration": mom_dict.get("duration", 0),
            "planned_start_time": mom_dict.get("planned_start_time", ""),
            "actual_start_time": mom_dict.get("actual_start_time", ""),
            "planned_end_time": mom_dict.get("planned_end_time", ""),
            "actual_end_time": mom_dict.get("actual_end_time", ""),
            "participants": mom_dict.get("participants", []),
            "introduction": mom_dict.get("introduction", ""),
            "points_discussed": mom_dict.get("points_discussed", []),
            "action_items": all_items,
            "conclusion": mom_dict.get("conclusion", ""),
        },
        "saved_at": dt_to_str(now),
        "label": "Regenerate: Action Points",
    }
    versions.append(new_version)

    async with get_db_context() as db:
        await db.execute(
            text("""
                UPDATE minutes_of_meeting SET
                    action_items = :action_items,
                    versions = :versions,
                    is_draft = 1,
                    updated_at = :updated_at
                WHERE recording_id = :rid AND user_id = :uid
            """),
            {
                "action_items": to_json(_normalize_action_items(all_items)),
                "versions": to_json(versions),
                "updated_at": dt_to_str(now),
                "rid": recording_id,
                "uid": user_id,
            },
        )
        await db.commit()

    logger.info(
        f"[MoM Regen Actions] Saved {len(all_items)} action item(s) as version #{new_version_num} "
        f"for recording_id={recording_id}"
    )

    return {
        "action_items": _normalize_action_items(all_items),
        "version_number": new_version_num,
        "windows_processed": total_windows,
    }


@router.post("/{recording_id}/summarize-long-points")
async def summarize_long_points_endpoint(
    recording_id: str,
    req: SummarizeLongPointsRequest,
    current_user: dict = Depends(get_current_user),
):
    """
    Summarize discussion points exceeding the given word threshold.
    Saves the updated MoM as a new version in history so the user can revert.
    """
    import asyncio as _asyncio

    user_id = current_user["id"]
    threshold = req.threshold

    async with get_db_context() as db:
        r = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        mom_row = r.mappings().fetchone()

    if not mom_row:
        raise HTTPException(status_code=404, detail="MoM not found")

    mom_dict = _mom_row_to_dict(mom_row)
    points = mom_dict.get("points_discussed", [])

    _loop = _asyncio.get_running_loop()
    try:
        updated_points, condensed_count = await _loop.run_in_executor(
            None,
            lambda: summarize_long_points(points, threshold=threshold),
        )
    finally:
        from services.ai_provider import QwenProvider
        QwenProvider.unload_model()

    now = datetime.now(timezone.utc)
    versions = from_json(mom_row["versions"], [])
    label = f"Summarize Long Points ({condensed_count} point{'s' if condensed_count != 1 else ''} > {threshold} words)"
    new_version_num = len(versions) + 1

    mom_dict["points_discussed"] = updated_points

    new_version = {
        "version": new_version_num,
        "data": {
            "title": mom_dict.get("title", ""),
            "date": mom_dict.get("date", ""),
            "duration": mom_dict.get("duration", 0),
            "planned_start_time": mom_dict.get("planned_start_time", ""),
            "actual_start_time": mom_dict.get("actual_start_time", ""),
            "planned_end_time": mom_dict.get("planned_end_time", ""),
            "actual_end_time": mom_dict.get("actual_end_time", ""),
            "participants": mom_dict.get("participants", []),
            "introduction": mom_dict.get("introduction", ""),
            "points_discussed": updated_points,
            "action_items": mom_dict.get("action_items", []),
            "conclusion": mom_dict.get("conclusion", ""),
        },
        "saved_at": dt_to_str(now),
        "label": label,
    }
    versions.append(new_version)

    async with get_db_context() as db:
        await db.execute(
            text("""
                UPDATE minutes_of_meeting SET
                    points_discussed = :points_discussed,
                    versions = :versions,
                    is_draft = 1,
                    updated_at = :updated_at
                WHERE recording_id = :rid AND user_id = :uid
            """),
            {
                "points_discussed": to_json(updated_points),
                "versions": to_json(versions),
                "updated_at": dt_to_str(now),
                "rid": recording_id,
                "uid": user_id,
            },
        )
        await db.commit()

        r2 = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        saved = r2.mappings().fetchone()

    saved_dict = _mom_row_to_dict(saved)

    return {
        "mom": saved_dict,
        "version_number": new_version_num,
        "condensed_count": condensed_count,
        "label": label,
    }




@router.post("/{recording_id}/rewrite/section")
async def rewrite_mom_section(
    recording_id: str,
    req: RewriteSectionRequest,
    current_user: dict = Depends(get_current_user),
):
    """
    Rewrite a MoM section (discussion_points, action_items, introduction, conclusion)
    according to provided style rules and custom prompt.

    For discussion_points and action_items, the content is processed in windows of
    req.window_size items to handle large sets.

    After rewriting, the MoM is saved and a new version entry is appended (regardless
    of the 5-minute debounce used by PATCH) with a label identifying the rewrite.
    Returns the full updated MoM.
    """
    import asyncio as _asyncio

    user_id = current_user["id"]
    section = req.section

    VALID_SECTIONS = {"discussion_points", "action_items", "introduction", "conclusion"}
    if section not in VALID_SECTIONS:
        raise HTTPException(status_code=400, detail=f"section must be one of {sorted(VALID_SECTIONS)}")

    logger.info(f"[MoM Rewrite Pipeline] Starting rewrite for recording_id='{recording_id}', section='{section}', user_id='{user_id}'")
    logger.info(f"[MoM Rewrite Pipeline] Request config: window_size={req.window_size}, rules_len={len(req.rules or '')}, custom_prompt_len={len(req.custom_prompt or '')}")

    # Fetch current MoM
    async with get_db_context() as db:
        r = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        mom_row = r.mappings().fetchone()

    if not mom_row:
        logger.error(f"[MoM Rewrite Pipeline ERROR] MoM not found for recording_id='{recording_id}', user_id='{user_id}'")
        raise HTTPException(status_code=404, detail="MoM not found")

    mom_dict = _mom_row_to_dict(mom_row)
    _loop = _asyncio.get_running_loop()

    try:
        if section in ("introduction", "conclusion"):
            # ── Single-pass rewrite for text fields ──────────────────
            original_text = str(req.content) if not isinstance(req.content, str) else req.content
            logger.info(f"[MoM Rewrite Step 1/6] Section '{section}' original text length={len(original_text)}:\n{original_text}")

            rewritten_text = await _loop.run_in_executor(
                None,
                lambda: rewrite_section_content(
                    content=original_text,
                    rules=req.rules,
                    custom_prompt=req.custom_prompt,
                    section_type=section,
                ),
            )
            rewritten_text = (rewritten_text or "").strip()
            if not rewritten_text:
                logger.error(f"[MoM Rewrite Step 2/6 ERROR] LLM returned empty result for {section}")
                raise HTTPException(status_code=500, detail=f"LLM generated an empty response for {section}. Rewrite aborted.")

            logger.info(f"[MoM Rewrite Step 3/6] Replacing '{section}' content:\nOLD:\n{original_text}\nNEW:\n{rewritten_text}")
            mom_dict[section] = rewritten_text

        else:
            # ── Windowed rewrite for list fields ──────────────────────
            if section == "discussion_points":
                if isinstance(req.content, list):
                    items = [str(x) for x in req.content]
                elif isinstance(req.content, str):
                    items = [x.strip() for x in req.content.splitlines() if x.strip()]
                else:
                    items = []
            else:  # action_items
                if isinstance(req.content, list):
                    items = []
                    for x in req.content:
                        if isinstance(x, dict):
                            items.append({
                                "task": str(x.get("task", "") or ""),
                                "owner": str(x.get("owner", "Unassigned") or "Unassigned"),
                                "deadline": str(x.get("deadline", "ASAP") or "ASAP"),
                            })
                        elif isinstance(x, str):
                            items.append(_parse_action_item_str(x))
                else:
                    items = mom_dict.get("action_items", [])

            if not items:
                logger.error(f"[MoM Rewrite Pipeline ERROR] No content items provided to rewrite for section='{section}'")
                raise HTTPException(status_code=400, detail=f"No items provided to rewrite for section '{section}'.")

            window = max(1, req.window_size)
            rewritten_items = []
            total_windows = (len(items) + window - 1) // window

            logger.info(f"[MoM Rewrite Step 1/6] List section '{section}' has {len(items)} items. Processing in {total_windows} window(s) of max size {window}.")

            for w_idx, start in enumerate(range(0, len(items), window), start=1):
                chunk = items[start: start + window]

                if section == "discussion_points":
                    chunk_text = "\n".join(f"{i+1}. {pt}" for i, pt in enumerate(chunk, start=start))
                else:  # action_items: send ONLY task description texts to LLM
                    chunk_tasks = [ai["task"] if isinstance(ai, dict) else str(ai) for ai in chunk]
                    chunk_text = "\n".join(f"{i+1}. {t}" for i, t in enumerate(chunk_tasks, start=start))

                logger.info(f"[MoM Rewrite Step 1/6] Window {w_idx}/{total_windows} (items {start+1}..{start+len(chunk)}): Input chunk text:\n{chunk_text}")

                try:
                    result_text = await _loop.run_in_executor(
                        None,
                        lambda ct=chunk_text: rewrite_section_content(
                            content=ct,
                            rules=req.rules,
                            custom_prompt=req.custom_prompt,
                            section_type=section,
                        ),
                    )
                except Exception as llm_err:
                    logger.error(f"[MoM Rewrite Step 1/6 ERROR] Window {w_idx}/{total_windows} LLM execution failed: {llm_err}", exc_info=True)
                    raise HTTPException(status_code=500, detail=f"LLM call failed for window {w_idx}/{total_windows}: {llm_err}")

                parsed = _parse_rewritten_list(result_text, section)
                logger.info(f"[MoM Rewrite Step 2/6] Window {w_idx}/{total_windows} Parsed Output ({len(parsed)} items):\n{json.dumps(parsed, indent=2)}")

                if not parsed:
                    logger.error(f"[MoM Rewrite Step 2/6 ERROR] Failed to parse any items from LLM response for window {w_idx}/{total_windows}. Raw output:\n{result_text}")
                    raise HTTPException(status_code=500, detail=f"Failed to parse LLM response for window {w_idx}/{total_windows}. Rewrite aborted.")

                if section == "discussion_points":
                    rewritten_items.extend(parsed)
                else:  # action_items: pair rewritten task texts with ORIGINAL owner and deadline!
                    for idx, task_text in enumerate(parsed):
                        if idx < len(chunk):
                            orig_owner = chunk[idx]["owner"] if isinstance(chunk[idx], dict) else "Unassigned"
                            orig_deadline = chunk[idx]["deadline"] if isinstance(chunk[idx], dict) else "ASAP"
                        else:
                            orig_owner = "Unassigned"
                            orig_deadline = "ASAP"
                        rewritten_items.append({
                            "task": task_text,
                            "owner": orig_owner,
                            "deadline": orig_deadline,
                        })

            if section == "discussion_points":
                orig_pts = mom_dict.get("points_discussed", [])
                logger.info(f"[MoM Rewrite Step 3/6] Replacing discussion_points (Old count={len(orig_pts)}, New count={len(rewritten_items)}):\nOLD:\n{json.dumps(orig_pts, indent=2)}\nNEW:\n{json.dumps(rewritten_items, indent=2)}")
                mom_dict["points_discussed"] = rewritten_items
            else:
                orig_ai = mom_dict.get("action_items", [])
                logger.info(f"[MoM Rewrite Step 3/6] Replacing action_items (Old count={len(orig_ai)}, New count={len(rewritten_items)}). PRESERVED original owners and deadlines:\nNEW ACTION ITEMS:\n{json.dumps(rewritten_items, indent=2)}")
                mom_dict["action_items"] = rewritten_items
    finally:
        from services.ai_provider import QwenProvider
        QwenProvider.unload_model()

    # ── Save to DB and push a new version entry ───────────────────────
    now = datetime.now(timezone.utc)
    versions = from_json(mom_row["versions"], [])
    rewrite_label = {
        "discussion_points": "Rewrite: Discussion Points",
        "action_items": "Rewrite: Action Points",
        "introduction": "Rewrite: Introduction",
        "conclusion": "Rewrite: Conclusion",
    }[section]

    new_version_num = len(versions) + 1
    new_version = {
        "version": new_version_num,
        "data": {
            "title": mom_dict.get("title", ""),
            "date": mom_dict.get("date", ""),
            "duration": mom_dict.get("duration", 0),
            "planned_start_time": mom_dict.get("planned_start_time", ""),
            "actual_start_time": mom_dict.get("actual_start_time", ""),
            "planned_end_time": mom_dict.get("planned_end_time", ""),
            "actual_end_time": mom_dict.get("actual_end_time", ""),
            "participants": mom_dict.get("participants", []),
            "introduction": mom_dict.get("introduction", ""),
            "points_discussed": mom_dict.get("points_discussed", []),
            "action_items": mom_dict.get("action_items", []),
            "conclusion": mom_dict.get("conclusion", ""),
        },
        "saved_at": dt_to_str(now),
        "label": rewrite_label,
    }
    versions.append(new_version)

    logger.info(f"[MoM Rewrite Step 4/6] Persisting updated MoM to SQLite DB. Creating Version #{new_version_num} ('{rewrite_label}')...")

    async with get_db_context() as db:
        await db.execute(
            text("""
                UPDATE minutes_of_meeting SET
                    title = :title, date = :date, duration = :duration,
                    planned_start_time = :planned_start_time,
                    actual_start_time = :actual_start_time,
                    planned_end_time = :planned_end_time,
                    actual_end_time = :actual_end_time,
                    participants = :participants,
                    introduction = :introduction,
                    points_discussed = :points_discussed,
                    action_items = :action_items,
                    conclusion = :conclusion,
                    versions = :versions, is_draft = 1, updated_at = :updated_at
                WHERE recording_id = :rid AND user_id = :uid
            """),
            {
                "title": mom_dict.get("title", ""),
                "date": mom_dict.get("date", ""),
                "duration": mom_dict.get("duration", 0),
                "planned_start_time": mom_dict.get("planned_start_time", ""),
                "actual_start_time": mom_dict.get("actual_start_time", ""),
                "planned_end_time": mom_dict.get("planned_end_time", ""),
                "actual_end_time": mom_dict.get("actual_end_time", ""),
                "participants": to_json(mom_dict.get("participants", [])),
                "introduction": mom_dict.get("introduction", ""),
                "points_discussed": to_json(mom_dict.get("points_discussed", [])),
                "action_items": to_json(_normalize_action_items(mom_dict.get("action_items", []))),
                "conclusion": mom_dict.get("conclusion", ""),
                "versions": to_json(versions),
                "updated_at": dt_to_str(now),
                "rid": recording_id,
                "uid": user_id,
            },
        )
        await db.commit()

        # Re-query saved record to verify DB write
        r2 = await db.execute(
            text("SELECT * FROM minutes_of_meeting WHERE recording_id = :rid AND user_id = :uid"),
            {"rid": recording_id, "uid": user_id},
        )
        saved = r2.mappings().fetchone()

    saved_dict = _mom_row_to_dict(saved)
    logger.info(f"[MoM Rewrite Step 4/6] DB save VERIFIED successfully!")
    logger.info(f"[MoM Rewrite Step 4/6] Verified DB points_discussed count={len(saved_dict.get('points_discussed', []))}:\n{json.dumps(saved_dict.get('points_discussed', []), indent=2)}")

    logger.info(f"[MoM Rewrite Step 5/6] Returning API response to frontend for section='{section}', version={new_version_num}")

    return {
        "mom": saved_dict,
        "version_number": new_version_num,
        "label": rewrite_label,
    }


# ── Helpers for parsing LLM-rewritten lists ──────────────────

def _parse_rewritten_list(text_output: str, section_type: str) -> List[str]:
    """
    Parse LLM-generated rewritten list output back into a list of strings.
    Handles numbered lists ("1. ...", "1) ..."), bullet lists ("- ...", "* ..."), and bare lines.
    Strips markdown code fences, headers, and preamble lines.
    """
    if not text_output or not text_output.strip():
        logger.warning(f"[MoM Rewrite Parsing] Received empty text_output for section_type={section_type}")
        return []

    lines = text_output.strip().splitlines()
    items = []

    preamble_starters = (
        "here is", "here are", "sure", "certainly", "below is", "below are",
        "revised list", "rewritten list", "updated list", "note:", "rules applied",
        "```markdown", "```json", "```text", "```", "written rules"
    )

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        # Skip markdown code fences
        if line.startswith("```"):
            continue

        line_lower = line.lower()

        # Skip preamble or header lines ending with ':' or starting with common LLM preamble
        if any(line_lower.startswith(p) for p in preamble_starters):
            if line.endswith(":") or len(line.split()) <= 8:
                logger.info(f"[MoM Rewrite Parsing] Filtered out LLM preamble line: {line!r}")
                continue

        # Strip common list prefixes: "1.", "1)", "-", "*", "•", "1 -", etc.
        cleaned = re.sub(r'^(\d+[\.\)\-]\s*|[\-*•]\s*)', '', line).strip()

        # Remove outer quotes if present
        if (cleaned.startswith('"') and cleaned.endswith('"')) or (cleaned.startswith("'") and cleaned.endswith("'")):
            cleaned = cleaned[1:-1].strip()

        if cleaned:
            items.append(cleaned)

    logger.info(f"[MoM Rewrite Parsing] Extracted {len(items)} items from {len(lines)} raw lines.")
    return items


def _parse_action_item_str(text: str) -> dict:
    """
    Parse an action item from the LLM's 'Task: X | Owner: Y | Deadline: Z' format.
    Falls back to treating the whole string as the task.
    """
    import re as _re
    if not isinstance(text, str):
        return {"task": str(text or ""), "owner": "Unassigned", "deadline": "ASAP"}

    text_clean = text.strip()

    # Try "Task: X | Owner: Y | Deadline: Z"
    m1 = _re.search(r'[Tt]ask:\s*(.+?)(?:\s*\|\s*[Oo]wner:\s*(.+?))?(?:\s*\|\s*[Dd]eadline:\s*(.+))?$', text_clean)
    if m1 and m1.group(1):
        task = m1.group(1).strip()
        owner = m1.group(2).strip() if m1.group(2) else "Unassigned"
        deadline = m1.group(3).strip() if m1.group(3) else "ASAP"
        return {"task": task, "owner": owner, "deadline": deadline}

    # Try "X (Owner: Y, Deadline: Z)"
    m2 = _re.search(r'^(.+?)\s*\([Oo]wner:\s*(.+?),\s*[Dd]eadline:\s*(.+?)\)$', text_clean)
    if m2:
        return {"task": m2.group(1).strip(), "owner": m2.group(2).strip(), "deadline": m2.group(3).strip()}

    return {"task": text_clean, "owner": "Unassigned", "deadline": "ASAP"}

